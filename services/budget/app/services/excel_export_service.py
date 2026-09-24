import io
from dataclasses import dataclass
from datetime import date, datetime, timezone
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.crud.budget_category_crud import list_budget_categories
from app.crud.budget_line_crud import list_budget_lines
from app.crud.currency_conversion_crud import FLOAT_EPSILON, list_currency_conversions
from app.crud.excel_export_crud import ReportLineExpense, get_report_line_expenses
from app.crud.funding_receipt_crud import list_funding_receipts
from app.models.budget import BudgetCategoryModel, BudgetLineModel, BudgetModel
from app.models.currency_ledger import CurrencyConversionModel, FundingReceiptModel
from app.services.budget_services import _add_duration_months, get_viewable_budget_service
from app.services.customer_client import CustomerServiceError, get_customer_cached
from app.services.user_cache import get_users_by_ids_cached

SHEET1_TITLE = "Original Budget"
SHEET2_TITLE = "Budget vs. Report Dashboard"
SHEET3_TITLE = "List of Expenses"
_BOLD = Font(bold=True)
_AUDIT_FONT = Font(italic=True, size=9, color="808080")
_TOTAL_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_TOP_BORDER = Border(top=Side(style="thin"))
_BOTTOM_BORDER = Border(bottom=Side(style="thin"))
_ESTIMATE_CELL_FONT = Font(italic=True)
_ESTIMATE_CELL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_FOOTNOTE_STAR_FONT = Font(italic=True, size=9, color="FF0000")
_ROW_FLAG_FONT = Font(color="FF0000")
_DESCRIPTION_COL_WIDTH = 26.63
_EXTRA_COL_WIDTH = 20.0
_AMOUNT_COL_WIDTH = 17.77
_ESTIMATE_COL_WIDTH = 17.53
_HEADER_ROW_COUNT = 6
_RATE_ROW = _HEADER_ROW_COUNT
_DESCRIPTION_COL = 1
_RATE_FORMAT = "0.0000"
_DATE_FORMAT = "yyyy-mm-dd"
_PERCENT_FORMAT = "0.0%"
_DASHBOARD_COL_WIDTHS = {1: 26.63, 2: 18.0, 3: 21.0, 4: 21.0, 5: 23.0, 6: 18.0}
_EXPENSE_DATE_COL_WIDTH = 12.0
_RATE_COL_WIDTH = 12.0
_EXPENSE_LIST_COL_WIDTHS = {
    1: _EXPENSE_DATE_COL_WIDTH,
    2: _DESCRIPTION_COL_WIDTH,
    3: _DESCRIPTION_COL_WIDTH,
    4: _DESCRIPTION_COL_WIDTH,
    5: _AMOUNT_COL_WIDTH,
    6: _ESTIMATE_COL_WIDTH,
    7: _RATE_COL_WIDTH,
    8: _EXPENSE_DATE_COL_WIDTH,
}


async def export_budget_workbook_service(
    db, valid_user: dict, budget_id: UUID
) -> tuple[BudgetModel, bytes]:
    """Loads one budget's categories/lines (auth: owner or funder, same as
    GET /budgets/{budget_id}) and builds its export workbook."""
    budget = await get_viewable_budget_service(budget_id, valid_user, db)
    categories = await list_budget_categories(db, budget_id=budget_id, limit=None)
    lines = await list_budget_lines(db, budget_id=budget_id, limit=None)
    conversions = await list_currency_conversions(db, budget_id=budget_id)
    receipts = await list_funding_receipts(db, budget_id=budget_id)
    expenses = await get_report_line_expenses(db, budget_id=budget_id)
    organisation_name = None
    try:
        organisation_name = (await get_customer_cached(budget.owner_id)).get("name")
    except CustomerServiceError:
        pass

    donor_name = budget.external_funder_name
    if budget.funding_customer_id:
        try:
            donor_name = (await get_customer_cached(budget.funding_customer_id)).get("name")
        except CustomerServiceError:
            donor_name = budget.external_funder_name

    try:
        users = await get_users_by_ids_cached(
            [str(valid_user["user_id"])], valid_user.get("token", "")
        )
    except Exception:
        users = {}
    exported_by = users.get(str(valid_user["user_id"]), {}).get("email")
    return budget, generate_budget_export_workbook(
        budget,
        categories,
        lines,
        conversions=conversions,
        receipts=receipts,
        expenses=expenses,
        organisation_name=organisation_name,
        donor_name=donor_name,
        exported_by=exported_by,
        exported_at=datetime.now(timezone.utc),
    )


def generate_budget_export_workbook(
    budget: BudgetModel,
    categories: list[BudgetCategoryModel],
    lines: list[BudgetLineModel],
    conversions: list[CurrencyConversionModel] | None = None,
    receipts: list[FundingReceiptModel] | None = None,
    expenses: list[ReportLineExpense] | None = None,
    organisation_name: str | None = None,
    donor_name: str | None = None,
    exported_by: str | None = None,
    exported_at: datetime | None = None,
) -> bytes:
    """Builds the export workbook for one budget. Group 6 makes sheet
    selection template-driven."""
    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = SHEET1_TITLE
    OriginalBudgetSheet(
        sheet1, budget, categories, lines, organisation_name, donor_name, exported_by, exported_at
    ).write()

    expense_rows = _build_expense_rows(
        lines, categories, expenses or [], budget.estimated_exchange_rate
    )

    sheet2 = wb.create_sheet(SHEET2_TITLE)
    DashboardSheet(
        sheet2,
        budget,
        organisation_name,
        donor_name,
        receipts or [],
        conversions or [],
        categories,
        lines,
        expense_rows,
        exported_by,
        exported_at,
    ).write()

    sheet3 = wb.create_sheet(SHEET3_TITLE)
    ExpenseListSheet(
        sheet3,
        budget,
        organisation_name,
        donor_name,
        expense_rows,
        exported_by,
        exported_at,
    ).write()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _period_label(start_date: date | None, duration_months: int | None) -> str | None:
    if not start_date:
        return None
    if not duration_months:  # None (unset) and 0 (BudgetModel's column default) both mean "unknown"
        return start_date.strftime("%m/%Y")
    end_date = _add_duration_months(start_date, duration_months - 1)
    return f"{start_date.strftime('%m/%Y')}-{end_date.strftime('%m/%Y')}"


@dataclass
class _ExpenseRow:
    expense_date: date
    budget_line_id: UUID | None
    budget_line_description: str | None
    category_name: str | None
    description: str | None
    amount: float
    rate: float | None
    is_estimated: bool
    conversion_date: date | None


def _usable_rate(rate: float | None) -> float | None:
    # A 0/negative rate would divide by zero or flip signs in Sheet 3's =E/G.
    return rate if rate and rate > 0 else None


def _implied_rate(allocation) -> float | None:
    if not allocation.conversion_donor_amount or not allocation.conversion_local_amount:
        return None
    return _usable_rate(allocation.conversion_local_amount / allocation.conversion_donor_amount)


def _build_expense_rows(
    lines: list[BudgetLineModel],
    categories: list[BudgetCategoryModel],
    expenses: list[ReportLineExpense],
    estimated_exchange_rate: float | None,
) -> list[_ExpenseRow]:
    """One row per allocation, plus one (estimated-rate-flagged) row per remainder."""
    lines_by_id = {line.id: line for line in lines}
    category_names = {category.id: category.name for category in categories}
    estimated_rate = _usable_rate(estimated_exchange_rate)
    rows: list[_ExpenseRow] = []
    for expense in expenses:
        budget_line = lines_by_id.get(expense.budget_line_id)
        budget_line_description = budget_line.description if budget_line else None
        category_name = category_names.get(budget_line.category_id) if budget_line else None
        allocated_total = 0.0
        for allocation in expense.allocations:
            allocated_total += allocation.amount_allocated
            rows.append(
                _ExpenseRow(
                    expense_date=expense.expense_date,
                    budget_line_id=expense.budget_line_id,
                    budget_line_description=budget_line_description,
                    category_name=category_name,
                    description=expense.description,
                    amount=allocation.amount_allocated,
                    rate=_implied_rate(allocation),
                    is_estimated=False,
                    conversion_date=allocation.converted_at,
                )
            )
        remainder = expense.amount - allocated_total
        # Rows must sum to expense.amount: keep refunds, zero amounts and over-allocation.
        if not expense.allocations or abs(remainder) > FLOAT_EPSILON:
            rows.append(
                _ExpenseRow(
                    expense_date=expense.expense_date,
                    budget_line_id=expense.budget_line_id,
                    budget_line_description=budget_line_description,
                    category_name=category_name,
                    description=expense.description,
                    amount=remainder,
                    rate=estimated_rate,
                    is_estimated=estimated_rate is not None,
                    conversion_date=None,
                )
            )
    return rows


def _group_lines_by_category(
    categories: list[BudgetCategoryModel], lines: list[BudgetLineModel]
) -> tuple[list[UUID | None], dict[UUID | None, str], dict[UUID | None, list[BudgetLineModel]]]:
    """Ordered category ids, display names, and lines — shared by Sheet 1/2."""
    categories_by_id: dict[UUID | None, BudgetCategoryModel] = {
        category.id: category for category in categories
    }
    lines_by_category_id: dict[UUID | None, list[BudgetLineModel]] = {}
    for line in lines:
        category = categories_by_id.get(line.category_id)
        lines_by_category_id.setdefault(category.id if category else None, []).append(line)

    # categories/lines already arrive ordered by created_at, id from the CRUD queries.
    ordered_category_ids: list[UUID | None] = list(categories_by_id)
    if None in lines_by_category_id:
        ordered_category_ids.append(None)

    category_names = {
        cid: (categories_by_id[cid].name if cid else "Uncategorized")
        for cid in ordered_category_ids
    }
    category_lines = {cid: lines_by_category_id.get(cid, []) for cid in ordered_category_ids}
    return ordered_category_ids, category_names, category_lines


def _audit_line(exported_by: str | None, exported_at: datetime | None) -> str:
    parts = ["Generated by OpenGrantFlow"]
    if exported_by:
        parts.append(exported_by)
    if exported_at:
        parts.append(exported_at.strftime("%Y-%m-%d %H:%M UTC"))
    return " · ".join(parts)


class _SheetWriter:
    """Shared cell-writing helpers for one worksheet (see design.md Decision 12)."""

    def __init__(
        self,
        ws,
        budget: BudgetModel,
        organisation_name: str | None,
        donor_name: str | None,
        exported_by: str | None = None,
        exported_at: datetime | None = None,
    ) -> None:
        self.ws = ws
        self.budget = budget
        self.organisation_name = organisation_name
        self.donor_name = donor_name
        self.exported_by = exported_by
        self.exported_at = exported_at

    def _write_header(self) -> str:
        """Writes the header block (rows 1-6) and returns the rate cell ref for formulas."""
        budget = self.budget
        fields = (
            ("Organisation Name", self.organisation_name),
            ("Donor Name", self.donor_name),
            ("Project Name", budget.name),
            ("Project Period", _period_label(budget.start_date, budget.duration_months)),
            ("Budget Original Currency", budget.actual_currency),
            ("Estimated Exchange Rate", budget.estimated_exchange_rate),
        )
        for row, (label, value) in enumerate(fields, start=1):
            self.ws.cell(row=row, column=1, value=label)
            self.ws.cell(row=row, column=2, value=value)
            self._bold_row(row)

        return f"$B${_RATE_ROW}"

    @staticmethod
    def _currency_format(currency: str | None) -> str:
        return f'#,##0.00" {currency}"' if currency else "#,##0.00"

    def _bold_row(self, row: int) -> None:
        for cell in self.ws[row]:
            cell.font = _BOLD

    def _apply_box_border(self, start_row: int, end_row: int, end_col: int) -> None:
        thin = Side(style="thin")
        for row in range(start_row, end_row + 1):
            for column in range(1, end_col + 1):
                self.ws.cell(row=row, column=column).border = Border(
                    top=thin if row == start_row else None,
                    bottom=thin if row == end_row else None,
                    left=thin if column == 1 else None,
                    right=thin if column == end_col else None,
                )

    def _set_cell(self, row: int, column: int, value, number_format: str):
        cell = self.ws.cell(row=row, column=column, value=value)
        cell.number_format = number_format
        return cell

    def _write_audit_footer(self, row: int) -> None:
        audit_line = _audit_line(self.exported_by, self.exported_at)
        cell = self.ws.cell(row=row, column=1, value=audit_line)
        cell.font = _AUDIT_FONT


class OriginalBudgetSheet(_SheetWriter):
    """Sheet 1 — Original Budget: header block, Budget Summary, Detailed Budget, footer."""

    def __init__(
        self,
        ws,
        budget: BudgetModel,
        categories: list[BudgetCategoryModel],
        lines: list[BudgetLineModel],
        organisation_name: str | None,
        donor_name: str | None,
        exported_by: str | None,
        exported_at: datetime | None,
    ) -> None:
        super().__init__(ws, budget, organisation_name, donor_name, exported_by, exported_at)
        self.categories = categories
        self.lines = lines

    def write(self) -> None:
        budget = self.budget
        has_rate = bool(budget.estimated_exchange_rate)
        rate_cell = self._write_header()

        local_fmt = self._currency_format(budget.local_currency)
        estimate_fmt = self._currency_format(budget.actual_currency)
        amount_header = f"Amount ({budget.local_currency})" if budget.local_currency else "Amount"
        estimate_header = (
            f"Estimate ({budget.actual_currency})" if budget.actual_currency else "Donor Estimate"
        )

        ordered_category_ids, category_names, category_lines = _group_lines_by_category(
            self.categories, self.lines
        )
        extra_keys = self._extra_field_keys()
        cols = self._amount_columns(extra_keys)
        self._set_column_widths(extra_keys, cols)

        plan = self._plan_rows(ordered_category_ids, category_lines)

        self._write_budget_summary(
            plan,
            ordered_category_ids,
            category_names,
            amount_header,
            estimate_header,
            extra_keys,
            cols,
            local_fmt,
            estimate_fmt,
            has_rate,
        )
        self._write_detailed_budget(
            plan,
            ordered_category_ids,
            category_names,
            category_lines,
            extra_keys,
            cols,
            local_fmt,
            estimate_fmt,
            has_rate,
            rate_cell,
        )
        self._write_footer(
            plan,
            ordered_category_ids,
            cols,
            local_fmt,
            estimate_fmt,
            has_rate,
        )

    def _extra_field_keys(self) -> list[str]:
        """Distinct extra_fields keys across the budget, in first-seen order."""
        keys: dict = {}
        for line in self.lines:
            for key in line.extra_fields or {}:
                keys[key] = None
        return list(keys)

    def _amount_columns(self, extra_keys: list[str]) -> dict:
        """Amount/Estimate shift right by one column per distinct extra field key."""
        amount_col = _DESCRIPTION_COL + 1 + len(extra_keys)
        estimate_col = amount_col + 1
        return {
            "amount_col": amount_col,
            "estimate_col": estimate_col,
            "amount_letter": get_column_letter(amount_col),
            "estimate_letter": get_column_letter(estimate_col),
        }

    def _plan_rows(self, ordered_category_ids: list, lines_by_category_id: dict) -> dict:
        """Computes every row number up front — Budget Summary formulas reference
        Detailed Budget's subtotal rows before those rows are written."""
        summary_header_row = _RATE_ROW + 3
        summary_rows = {}
        row = summary_header_row
        for category_id in ordered_category_ids:
            row += 1
            summary_rows[category_id] = row
        summary_total_row = row + 1 if ordered_category_ids else summary_header_row + 1

        detail_title_row = summary_total_row + 2
        detail_header_rows: dict = {}
        detail_line_rows: dict = {}
        detail_subtotal_rows: dict = {}
        row = detail_title_row
        for category_id in ordered_category_ids:
            row += 1
            detail_header_rows[category_id] = row
            line_rows = []
            for _ in lines_by_category_id.get(category_id, []):
                row += 1
                line_rows.append(row)
            detail_line_rows[category_id] = line_rows
            row += 1
            detail_subtotal_rows[category_id] = row
            row += 1  # blank separator after the category block

        footer_total_row = row + 2 if ordered_category_ids else detail_title_row + 2
        signature_line_row = footer_total_row + 3
        contact_line_row = signature_line_row + 4
        audit_row = contact_line_row + 3

        return {
            "summary_header_row": summary_header_row,
            "summary_rows": summary_rows,
            "summary_total_row": summary_total_row,
            "detail_title_row": detail_title_row,
            "detail_header_rows": detail_header_rows,
            "detail_line_rows": detail_line_rows,
            "detail_subtotal_rows": detail_subtotal_rows,
            "footer_total_row": footer_total_row,
            "signature_line_row": signature_line_row,
            "contact_line_row": contact_line_row,
            "audit_row": audit_row,
        }

    def _set_column_widths(self, extra_keys: list[str], cols: dict) -> None:
        self.ws.column_dimensions[get_column_letter(_DESCRIPTION_COL)].width = (
            _DESCRIPTION_COL_WIDTH
        )
        for i in range(len(extra_keys)):
            self.ws.column_dimensions[get_column_letter(_DESCRIPTION_COL + 1 + i)].width = (
                _EXTRA_COL_WIDTH
            )
        self.ws.column_dimensions[get_column_letter(cols["amount_col"])].width = _AMOUNT_COL_WIDTH
        self.ws.column_dimensions[get_column_letter(cols["estimate_col"])].width = (
            _ESTIMATE_COL_WIDTH
        )

    def _write_budget_summary(
        self,
        plan,
        ordered_category_ids,
        category_names,
        amount_header,
        estimate_header,
        extra_keys,
        cols,
        local_fmt,
        estimate_fmt,
        has_rate,
    ) -> None:
        ws = self.ws
        amount_col, estimate_col = cols["amount_col"], cols["estimate_col"]
        amount_letter, estimate_letter = cols["amount_letter"], cols["estimate_letter"]

        header_row = plan["summary_header_row"]
        ws.cell(row=header_row, column=1, value="BUDGET SUMMARY")
        for i, key in enumerate(extra_keys):
            ws.cell(row=header_row, column=_DESCRIPTION_COL + 1 + i, value=key)
        ws.cell(row=header_row, column=amount_col, value=amount_header)
        ws.cell(row=header_row, column=estimate_col, value=estimate_header)
        self._bold_row(header_row)

        for category_id in ordered_category_ids:
            row = plan["summary_rows"][category_id]
            detail_subtotal_row = plan["detail_subtotal_rows"][category_id]
            ws.cell(row=row, column=1, value=category_names[category_id])
            self._set_cell(row, amount_col, f"={amount_letter}{detail_subtotal_row}", local_fmt)
            if has_rate:
                self._set_cell(
                    row, estimate_col, f"={estimate_letter}{detail_subtotal_row}", estimate_fmt
                )
        if ordered_category_ids:
            first_summary_row = plan["summary_rows"][ordered_category_ids[0]]
            last_summary_row = plan["summary_rows"][ordered_category_ids[-1]]
            self._apply_box_border(first_summary_row, last_summary_row, estimate_col)

        total_row = plan["summary_total_row"]
        ws.cell(row=total_row, column=1, value="TOTAL")
        if ordered_category_ids:
            first_row = plan["summary_rows"][ordered_category_ids[0]]
            last_row = plan["summary_rows"][ordered_category_ids[-1]]
            amount_range = f"{amount_letter}{first_row}:{amount_letter}{last_row}"
            self._set_cell(total_row, amount_col, f"=SUM({amount_range})", local_fmt)
            if has_rate:
                estimate_range = f"{estimate_letter}{first_row}:{estimate_letter}{last_row}"
                self._set_cell(total_row, estimate_col, f"=SUM({estimate_range})", estimate_fmt)
        else:
            self._set_cell(total_row, amount_col, 0.0, local_fmt)
        self._bold_row(total_row)

    def _write_detailed_budget(
        self,
        plan,
        ordered_category_ids,
        category_names,
        category_lines,
        extra_keys,
        cols,
        local_fmt,
        estimate_fmt,
        has_rate,
        rate_cell,
    ) -> None:
        ws = self.ws
        amount_col, estimate_col = cols["amount_col"], cols["estimate_col"]
        amount_letter, estimate_letter = cols["amount_letter"], cols["estimate_letter"]

        ws.cell(row=plan["detail_title_row"], column=1, value="DETAILED BUDGET")
        self._bold_row(plan["detail_title_row"])

        for category_id in ordered_category_ids:
            header_row = plan["detail_header_rows"][category_id]
            ws.cell(row=header_row, column=1, value=category_names[category_id])
            self._bold_row(header_row)

            line_rows = plan["detail_line_rows"][category_id]
            for line, row in zip(category_lines[category_id], line_rows):
                ws.cell(row=row, column=1, value=line.description)
                for i, key in enumerate(extra_keys):
                    value = (line.extra_fields or {}).get(key)
                    ws.cell(row=row, column=_DESCRIPTION_COL + 1 + i, value=value)
                self._set_cell(row, amount_col, line.amount or 0.0, local_fmt)
                if has_rate:
                    self._set_cell(
                        row, estimate_col, f"={amount_letter}{row}/{rate_cell}", estimate_fmt
                    )
            if line_rows:
                self._apply_box_border(line_rows[0], line_rows[-1], estimate_col)

            subtotal_row = plan["detail_subtotal_rows"][category_id]
            ws.cell(row=subtotal_row, column=1, value="Subtotal")
            if line_rows:
                amount_range = f"{amount_letter}{line_rows[0]}:{amount_letter}{line_rows[-1]}"
                self._set_cell(subtotal_row, amount_col, f"=SUM({amount_range})", local_fmt)
                if has_rate:
                    estimate_range = (
                        f"{estimate_letter}{line_rows[0]}:{estimate_letter}{line_rows[-1]}"
                    )
                    self._set_cell(
                        subtotal_row, estimate_col, f"=SUM({estimate_range})", estimate_fmt
                    )
            else:
                self._set_cell(subtotal_row, amount_col, 0.0, local_fmt)
                if has_rate:
                    formula = f"={amount_letter}{subtotal_row}/{rate_cell}"
                    self._set_cell(subtotal_row, estimate_col, formula, estimate_fmt)
            self._bold_row(subtotal_row)

    def _write_footer(
        self,
        plan,
        ordered_category_ids,
        cols,
        local_fmt,
        estimate_fmt,
        has_rate,
    ) -> None:
        ws = self.ws
        amount_col, estimate_col = cols["amount_col"], cols["estimate_col"]
        amount_letter, estimate_letter = cols["amount_letter"], cols["estimate_letter"]

        total_row = plan["footer_total_row"]
        ws.cell(row=total_row, column=1, value="Total expenditures")
        if ordered_category_ids:
            subtotal_rows = [plan["detail_subtotal_rows"][cid] for cid in ordered_category_ids]
            c_formula = "=" + "+".join(f"{amount_letter}{r}" for r in subtotal_rows)
            self._set_cell(total_row, amount_col, c_formula, local_fmt)
            if has_rate:
                d_formula = "=" + "+".join(f"{estimate_letter}{r}" for r in subtotal_rows)
                self._set_cell(total_row, estimate_col, d_formula, estimate_fmt)
        else:
            self._set_cell(total_row, amount_col, 0.0, local_fmt)
        for cell in ws[total_row]:
            cell.font = _BOLD
            cell.fill = _TOTAL_FILL

        signature_row = plan["signature_line_row"]
        ws.cell(row=signature_row, column=1).border = _TOP_BORDER
        ws.cell(row=signature_row + 1, column=1, value="Authorised Signatory")

        contact_row = plan["contact_line_row"]
        ws.cell(row=contact_row, column=1).border = _TOP_BORDER
        ws.cell(row=contact_row + 1, column=1, value="Project contact person")

        self._write_audit_footer(plan["audit_row"])


class DashboardSheet(_SheetWriter):
    """Sheet 2 — Budget vs. Report Dashboard: donor-currency-first framing."""

    def __init__(
        self,
        ws,
        budget: BudgetModel,
        organisation_name: str | None,
        donor_name: str | None,
        receipts: list[FundingReceiptModel],
        conversions: list[CurrencyConversionModel],
        categories: list[BudgetCategoryModel],
        lines: list[BudgetLineModel],
        expense_rows: list["_ExpenseRow"],
        exported_by: str | None,
        exported_at: datetime | None,
    ) -> None:
        super().__init__(ws, budget, organisation_name, donor_name, exported_by, exported_at)
        self.receipts = receipts
        self.conversions = conversions
        self.categories = categories
        self.lines = lines
        self.expense_rows = expense_rows

    def write(self) -> None:
        self._set_column_widths()
        self._write_header()
        ordered_category_ids, category_names, category_lines = _group_lines_by_category(
            self.categories, self.lines
        )
        events = self._ledger_events()
        plan = self._plan(ordered_category_ids, category_lines, len(events))
        estimated_exchange_rate = self.budget.estimated_exchange_rate

        approved_total = (
            sum(line.amount or 0.0 for line in self.lines) / estimated_exchange_rate
            if estimated_exchange_rate
            else None
        )
        received_total = sum(receipt.amount for receipt in self.receipts)
        converted_total = sum(conversion.donor_amount for conversion in self.conversions)
        local_converted_total = sum(conversion.local_amount for conversion in self.conversions)
        local_expenses_total = sum(row.amount for row in self.expense_rows)

        self._write_approved_block(plan, approved_total)
        self._write_balance_block(
            plan,
            approved_total,
            received_total,
            converted_total,
            local_converted_total - local_expenses_total,
        )
        self._write_ledger(plan, events)
        self._write_report_summary(plan, ordered_category_ids, category_names)
        estimated_line_ids = {row.budget_line_id for row in self.expense_rows if row.is_estimated}
        has_estimated_cells = self._write_detail(
            plan,
            ordered_category_ids,
            category_names,
            category_lines,
            bool(estimated_exchange_rate),
            estimated_line_ids,
        )
        self._write_footer(plan, has_estimated_cells)

    def _set_column_widths(self) -> None:
        for column, width in _DASHBOARD_COL_WIDTHS.items():
            self.ws.column_dimensions[get_column_letter(column)].width = width
        self.ws.column_dimensions[get_column_letter(8)].hidden = True

    def _ledger_events(
        self,
    ) -> list[tuple[date, int, FundingReceiptModel | CurrencyConversionModel]]:
        """Receipts and conversions merged into one date-ordered ledger; a
        receipt sorts before a conversion sharing its date."""
        events: list[tuple[date, int, FundingReceiptModel | CurrencyConversionModel]] = [
            (receipt.received_at, 0, receipt) for receipt in self.receipts
        ]
        events += [(conversion.converted_at, 1, conversion) for conversion in self.conversions]
        events.sort(key=lambda event: (event[0], event[1]))
        return events

    def _plan(self, ordered_category_ids, category_lines, ledger_event_count: int) -> dict:
        approved_total_row = _HEADER_ROW_COUNT + 2
        approved_on_row = approved_total_row + 1
        received_total_row = approved_total_row + 3
        converted_total_row = received_total_row + 1
        balance_header_row = received_total_row + 2
        balance_value_row = received_total_row + 3

        ledger_title_row = balance_value_row + 2
        ledger_header_row = ledger_title_row + 1
        ledger_data_start_row = ledger_header_row + 1
        ledger_total_row = (
            ledger_data_start_row + ledger_event_count
            if ledger_event_count
            else ledger_header_row + 1
        )

        report_summary_header_row = ledger_total_row + 3
        report_summary_rows: dict = {}
        row = report_summary_header_row
        for category_id in ordered_category_ids:
            row += 1
            report_summary_rows[category_id] = row
        report_summary_total_row = (
            row + 1 if ordered_category_ids else report_summary_header_row + 1
        )

        detail_category_header_rows: dict = {}
        detail_line_rows: dict = {}
        detail_subtotal_rows: dict = {}
        row = report_summary_total_row + 1
        for category_id in ordered_category_ids:
            row += 1
            detail_category_header_rows[category_id] = row
            line_rows = []
            for _ in category_lines[category_id]:
                row += 1
                line_rows.append(row)
            detail_line_rows[category_id] = line_rows
            subtotal_row = row + 1
            detail_subtotal_rows[category_id] = subtotal_row
            row = subtotal_row + 1  # blank separator row before the next category

        refund_row = row + 4
        place_date_row = refund_row + 4
        signature_line_row = place_date_row + 4
        signature_label_row = signature_line_row + 1
        audit_row = signature_label_row + 4

        return {
            "approved_total_row": approved_total_row,
            "approved_on_row": approved_on_row,
            "received_total_row": received_total_row,
            "converted_total_row": converted_total_row,
            "balance_header_row": balance_header_row,
            "balance_value_row": balance_value_row,
            "ledger_title_row": ledger_title_row,
            "ledger_header_row": ledger_header_row,
            "ledger_data_start_row": ledger_data_start_row,
            "ledger_total_row": ledger_total_row,
            "report_summary_header_row": report_summary_header_row,
            "report_summary_rows": report_summary_rows,
            "report_summary_total_row": report_summary_total_row,
            "detail_category_header_rows": detail_category_header_rows,
            "detail_line_rows": detail_line_rows,
            "detail_subtotal_rows": detail_subtotal_rows,
            "refund_row": refund_row,
            "place_date_row": place_date_row,
            "signature_line_row": signature_line_row,
            "signature_label_row": signature_label_row,
            "audit_row": audit_row,
        }

    def _write_approved_block(self, plan: dict, approved_total: float | None) -> None:
        ws = self.ws
        donor_fmt = self._currency_format(self.budget.actual_currency)

        ws.cell(row=plan["approved_total_row"], column=1, value="Approved Total")
        if approved_total is not None:
            self._set_cell(plan["approved_total_row"], 2, approved_total, donor_fmt)

        ws.cell(row=plan["approved_on_row"], column=1, value="Approved on")
        if self.budget.confirmed_at:
            # openpyxl rejects tz-aware datetimes outright; Excel has no timezone concept.
            confirmed_at = self.budget.confirmed_at.replace(tzinfo=None)
            cell = ws.cell(row=plan["approved_on_row"], column=2, value=confirmed_at)
            cell.number_format = _DATE_FORMAT

    def _write_balance_block(
        self,
        plan: dict,
        approved_total: float | None,
        received_total: float,
        converted_total: float,
        local_balance: float,
    ) -> None:
        ws = self.ws
        donor_fmt = self._currency_format(self.budget.actual_currency)
        local_fmt = self._currency_format(self.budget.local_currency)

        received_row = plan["received_total_row"]
        ws.cell(row=received_row, column=1, value="Received Total")
        self._set_cell(received_row, 2, received_total, donor_fmt)
        if approved_total:
            self._set_cell(
                received_row, 3, f"=B{received_row}/B{plan['approved_total_row']}", _PERCENT_FORMAT
            )

        converted_row = plan["converted_total_row"]
        ws.cell(row=converted_row, column=1, value="Converted Total")
        self._set_cell(converted_row, 2, converted_total, donor_fmt)
        if received_total:
            self._set_cell(converted_row, 3, f"=B{converted_row}/B{received_row}", _PERCENT_FORMAT)

        header_row = plan["balance_header_row"]
        ws.cell(row=header_row, column=1, value="Current Balance")
        ws.cell(row=header_row, column=2, value=self.budget.actual_currency or "Donor")
        ws.cell(row=header_row, column=3, value=self.budget.local_currency or "Local")
        self._bold_row(header_row)

        value_row = plan["balance_value_row"]
        self._set_cell(value_row, 2, f"=B{received_row}-B{converted_row}", donor_fmt)
        self._set_cell(value_row, 3, local_balance, local_fmt)

    def _write_ledger(
        self,
        plan: dict,
        events: list[tuple[date, int, FundingReceiptModel | CurrencyConversionModel]],
    ) -> None:
        ws = self.ws
        budget = self.budget
        donor_fmt = self._currency_format(budget.actual_currency)
        local_fmt = self._currency_format(budget.local_currency)
        donor_header = (
            f"Donor Amount ({budget.actual_currency})" if budget.actual_currency else "Donor Amount"
        )
        local_header = (
            f"Local Amount ({budget.local_currency})" if budget.local_currency else "Local Amount"
        )

        ws.cell(row=plan["ledger_title_row"], column=1, value="Funding Ledger")
        self._bold_row(plan["ledger_title_row"])

        header_row = plan["ledger_header_row"]
        ws.cell(row=header_row, column=1, value="Date")
        ws.cell(row=header_row, column=2, value="Received")
        ws.cell(row=header_row, column=3, value=donor_header)
        ws.cell(row=header_row, column=4, value=local_header)
        ws.cell(row=header_row, column=5, value="Implied Rate")
        self._bold_row(header_row)

        row = plan["ledger_data_start_row"] - 1
        for event_date, kind, obj in events:
            row += 1
            ws.cell(row=row, column=1, value=event_date).number_format = _DATE_FORMAT
            if isinstance(obj, FundingReceiptModel):
                self._set_cell(row, 2, obj.amount, donor_fmt)
            else:
                self._set_cell(row, 3, obj.donor_amount, donor_fmt)
                self._set_cell(row, 4, obj.local_amount, local_fmt)
                if obj.donor_amount:
                    self._set_cell(row, 5, f"=D{row}/C{row}", _RATE_FORMAT)

        total_row = plan["ledger_total_row"]
        ws.cell(row=total_row, column=1, value="TOTAL")
        if events:
            first_row, last_row = plan["ledger_data_start_row"], row
            self._set_cell(total_row, 3, f"=SUM(C{first_row}:C{last_row})", donor_fmt)
            self._set_cell(total_row, 4, f"=SUM(D{first_row}:D{last_row})", local_fmt)
        else:
            self._set_cell(total_row, 3, 0.0, donor_fmt)
            self._set_cell(total_row, 4, 0.0, local_fmt)
        self._bold_row(total_row)

    def _write_report_summary(self, plan: dict, ordered_category_ids, category_names) -> None:
        ws = self.ws
        budget = self.budget
        donor_fmt = self._currency_format(budget.actual_currency)
        local_fmt = self._currency_format(budget.local_currency)

        header_row = plan["report_summary_header_row"]
        ws.cell(row=header_row, column=1, value="Report Summary")
        ws.cell(
            row=header_row, column=2, value=self._column_header("Original", budget.actual_currency)
        )
        ws.cell(
            row=header_row, column=3, value=self._column_header("Planned", budget.local_currency)
        )
        ws.cell(
            row=header_row, column=4, value=self._column_header("Expenses", budget.local_currency)
        )
        ws.cell(
            row=header_row,
            column=5,
            value=self._column_header("Total Expenses", budget.actual_currency),
        )
        ws.cell(
            row=header_row, column=6, value=self._column_header("Deviation", budget.actual_currency)
        )
        self._bold_row(header_row)

        for category_id in ordered_category_ids:
            row = plan["report_summary_rows"][category_id]
            subtotal_row = plan["detail_subtotal_rows"][category_id]
            ws.cell(row=row, column=1, value=category_names[category_id])
            self._set_cell(row, 2, f"=B{subtotal_row}", donor_fmt)
            self._set_cell(row, 3, f"=C{subtotal_row}", local_fmt)
            self._set_cell(row, 4, f"=D{subtotal_row}", local_fmt)
            self._set_cell(row, 5, f"=E{subtotal_row}", donor_fmt)
            self._set_cell(row, 6, f"=F{subtotal_row}", donor_fmt)
        if ordered_category_ids:
            first_summary_row = plan["report_summary_rows"][ordered_category_ids[0]]
            last_summary_row = plan["report_summary_rows"][ordered_category_ids[-1]]
            self._apply_box_border(first_summary_row, last_summary_row, 6)

        total_row = plan["report_summary_total_row"]
        ws.cell(row=total_row, column=1, value="TOTAL")
        column_formats = (
            (2, donor_fmt),
            (3, local_fmt),
            (4, local_fmt),
            (5, donor_fmt),
            (6, donor_fmt),
        )
        if ordered_category_ids:
            first_row = plan["report_summary_rows"][ordered_category_ids[0]]
            last_row = plan["report_summary_rows"][ordered_category_ids[-1]]
            for column, fmt in column_formats:
                letter = get_column_letter(column)
                self._set_cell(
                    total_row, column, f"=SUM({letter}{first_row}:{letter}{last_row})", fmt
                )
        else:
            for column, fmt in column_formats:
                self._set_cell(total_row, column, 0.0, fmt)
        self._bold_row(total_row)
        self._apply_box_border(total_row, total_row, 6)

    def _write_detail(
        self,
        plan: dict,
        ordered_category_ids,
        category_names,
        category_lines,
        has_rate: bool,
        estimated_line_ids: set,
    ) -> bool:
        ws = self.ws
        budget = self.budget
        donor_fmt = self._currency_format(budget.actual_currency)
        local_fmt = self._currency_format(budget.local_currency)
        estimated_exchange_rate = budget.estimated_exchange_rate
        has_estimated_cells = False

        for category_id in ordered_category_ids:
            header_row = plan["detail_category_header_rows"][category_id]
            ws.cell(row=header_row, column=1, value=category_names[category_id])
            self._bold_row(header_row)

            line_rows = plan["detail_line_rows"][category_id]
            for line, row in zip(category_lines[category_id], line_rows):
                planned = line.amount or 0.0
                ws.cell(row=row, column=1, value=line.description)
                self._set_cell(row, 3, planned, local_fmt)
                ws.cell(row=row, column=8, value=str(line.id))
                self._set_cell(
                    row,
                    4,
                    f"=SUMIF('{SHEET3_TITLE}'!$I:$I,$H{row},'{SHEET3_TITLE}'!$E:$E)",
                    local_fmt,
                )
                if has_rate and estimated_exchange_rate:
                    self._set_cell(row, 2, planned / estimated_exchange_rate, donor_fmt)
                    cell = self._set_cell(
                        row,
                        5,
                        f"=SUMIF('{SHEET3_TITLE}'!$I:$I,$H{row},'{SHEET3_TITLE}'!$F:$F)",
                        donor_fmt,
                    )
                    if line.id in estimated_line_ids:
                        cell.font = _ESTIMATE_CELL_FONT
                        cell.fill = _ESTIMATE_CELL_FILL
                        ws.cell(row=row, column=7, value="*").font = _ROW_FLAG_FONT
                        has_estimated_cells = True
                    self._set_cell(row, 6, f"=B{row}-E{row}", donor_fmt)
            if line_rows:
                self._apply_box_border(line_rows[0], line_rows[-1], 6)

            subtotal_row = plan["detail_subtotal_rows"][category_id]
            ws.cell(row=subtotal_row, column=1, value="Subtotal")
            if line_rows:
                first_row, last_row = line_rows[0], line_rows[-1]
                self._set_cell(subtotal_row, 3, f"=SUM(C{first_row}:C{last_row})", local_fmt)
                self._set_cell(subtotal_row, 4, f"=SUM(D{first_row}:D{last_row})", local_fmt)
                if has_rate:
                    self._set_cell(subtotal_row, 2, f"=SUM(B{first_row}:B{last_row})", donor_fmt)
                    self._set_cell(subtotal_row, 5, f"=SUM(E{first_row}:E{last_row})", donor_fmt)
                    self._set_cell(subtotal_row, 6, f"=SUM(F{first_row}:F{last_row})", donor_fmt)
            else:
                self._set_cell(subtotal_row, 3, 0.0, local_fmt)
                self._set_cell(subtotal_row, 4, 0.0, local_fmt)
                if has_rate:
                    self._set_cell(subtotal_row, 2, 0.0, donor_fmt)
                    self._set_cell(subtotal_row, 5, 0.0, donor_fmt)
                    self._set_cell(subtotal_row, 6, f"=B{subtotal_row}-E{subtotal_row}", donor_fmt)
            self._bold_row(subtotal_row)
            self._apply_box_border(subtotal_row, subtotal_row, 6)

        return has_estimated_cells

    def _write_footer(self, plan: dict, has_estimated_cells: bool = False) -> None:
        ws = self.ws
        donor_fmt = self._currency_format(self.budget.actual_currency)

        ws.cell(row=plan["refund_row"], column=1, value="Refund to donor:")
        refund_cell = self._set_cell(
            plan["refund_row"],
            2,
            f"=B{plan['received_total_row']}-E{plan['report_summary_total_row']}",
            donor_fmt,
        )
        refund_cell.font = _BOLD

        ws.cell(row=plan["place_date_row"], column=1, value="Place, date:")
        ws.cell(row=plan["place_date_row"], column=2).border = _BOTTOM_BORDER

        line_row = plan["signature_line_row"]
        for column in (1, 2, 4, 5):
            ws.cell(row=line_row, column=column).border = _BOTTOM_BORDER

        label_row = plan["signature_label_row"]
        ws.cell(row=label_row, column=1, value="Authorised Signatory")
        ws.cell(row=label_row, column=4, value="Project Contact Person")

        if has_estimated_cells:
            legend_row = plan["audit_row"] - 2
            ws.cell(row=legend_row, column=1, value="*").font = _FOOTNOTE_STAR_FONT
            legend_cell = ws.cell(
                row=legend_row,
                column=2,
                value=(
                    "Highlighted amount includes an estimate — the budget's estimated "
                    "exchange rate applied to the portion not yet linked to a currency conversion."
                ),
            )
            legend_cell.font = _AUDIT_FONT

        self._write_audit_footer(plan["audit_row"])

    @staticmethod
    def _column_header(label: str, currency: str | None) -> str:
        return f"{label} ({currency})" if currency else label


class ExpenseListSheet(_SheetWriter):
    """Sheet 3 — List of Expenses: one row per expense, or per allocation subline (Decision 12)."""

    def __init__(
        self,
        ws,
        budget: BudgetModel,
        organisation_name: str | None,
        donor_name: str | None,
        rows: list["_ExpenseRow"],
        exported_by: str | None,
        exported_at: datetime | None,
    ) -> None:
        super().__init__(ws, budget, organisation_name, donor_name, exported_by, exported_at)
        self.rows = rows

    def write(self) -> None:
        self._set_column_widths()
        rows = self.rows
        needs_footnote = self._needs_rate_footnote(rows)
        plan = self._plan(len(rows), needs_footnote)
        self._write_table(plan, rows, needs_footnote)
        self._write_footer(plan)

    def _set_column_widths(self) -> None:
        for column, width in _EXPENSE_LIST_COL_WIDTHS.items():
            self.ws.column_dimensions[get_column_letter(column)].width = width
        self.ws.column_dimensions[get_column_letter(9)].hidden = True

    @staticmethod
    def _needs_rate_footnote(rows: list["_ExpenseRow"]) -> bool:
        """True whenever a row is estimated or lacks a rate entirely."""
        return any(row.is_estimated or row.rate is None for row in rows)

    def _plan(self, row_count: int, needs_footnote: bool = False) -> dict:
        header_row = 1
        data_start_row = header_row + 1
        last_row = data_start_row + row_count - 1 if row_count else header_row
        total_row = last_row + 1 if row_count else header_row + 1
        if needs_footnote:
            footnote_row = total_row + 2
            audit_row = footnote_row + 2
        else:
            footnote_row = None
            audit_row = total_row + 2
        return {
            "header_row": header_row,
            "data_start_row": data_start_row,
            "last_row": last_row,
            "total_row": total_row,
            "footnote_row": footnote_row,
            "audit_row": audit_row,
        }

    def _write_table(self, plan: dict, rows: list["_ExpenseRow"], needs_footnote: bool) -> None:
        ws = self.ws
        budget = self.budget
        local_fmt = self._currency_format(budget.local_currency)
        donor_fmt = self._currency_format(budget.actual_currency)

        header_row = plan["header_row"]
        ws.cell(row=header_row, column=1, value="Date")
        ws.cell(row=header_row, column=2, value="Category")
        ws.cell(row=header_row, column=3, value="Budget Line")
        ws.cell(row=header_row, column=4, value="Description")
        ws.cell(
            row=header_row,
            column=5,
            value=f"Amount ({budget.local_currency})" if budget.local_currency else "Amount",
        )
        ws.cell(
            row=header_row,
            column=6,
            value=(
                f"Converted Amount ({budget.actual_currency})"
                if budget.actual_currency
                else "Converted Amount"
            ),
        )
        ws.cell(row=header_row, column=7, value="Exchange rate")
        ws.cell(row=header_row, column=8, value="Exchange date")
        ws.cell(row=header_row, column=9, value="Budget Line ID")
        self._bold_row(header_row)

        row = plan["data_start_row"] - 1
        for expense_row in rows:
            row += 1
            ws.cell(row=row, column=1, value=expense_row.expense_date).number_format = _DATE_FORMAT
            ws.cell(row=row, column=2, value=expense_row.category_name)
            ws.cell(row=row, column=3, value=expense_row.budget_line_description)
            ws.cell(row=row, column=4, value=expense_row.description)
            self._set_cell(row, 5, expense_row.amount, local_fmt)
            if expense_row.rate is not None:
                self._set_cell(row, 7, expense_row.rate, _RATE_FORMAT)
                self._set_cell(row, 6, f"=E{row}/G{row}", donor_fmt)
            if expense_row.conversion_date is not None:
                ws.cell(row=row, column=8, value=expense_row.conversion_date).number_format = (
                    _DATE_FORMAT
                )
            elif expense_row.is_estimated or expense_row.rate is None:
                ws.cell(row=row, column=8, value="*").font = _ROW_FLAG_FONT
            ws.cell(
                row=row,
                column=9,
                value=str(expense_row.budget_line_id) if expense_row.budget_line_id else None,
            )

        if rows:
            self._apply_box_border(plan["data_start_row"], plan["last_row"], 8)

        total_row = plan["total_row"]
        ws.cell(row=total_row, column=1, value="Total")
        if rows:
            first_row, last_row = plan["data_start_row"], plan["last_row"]
            self._set_cell(total_row, 5, f"=SUM(E{first_row}:E{last_row})", local_fmt)
            if any(expense_row.rate is not None for expense_row in rows):
                self._set_cell(total_row, 6, f"=SUM(F{first_row}:F{last_row})", donor_fmt)
                # A blended rate is only meaningful when every row has one; otherwise
                # it'd divide all-local (E) by only-rated (F), fabricating a rate.
                if all(expense_row.rate is not None for expense_row in rows):
                    self._set_cell(total_row, 7, f"=E{total_row}/F{total_row}", _RATE_FORMAT)
        else:
            self._set_cell(total_row, 5, 0.0, local_fmt)
        self._bold_row(total_row)

        if needs_footnote:
            estimated = sum(1 for expense_row in rows if expense_row.is_estimated)
            no_rate = sum(1 for expense_row in rows if expense_row.rate is None)
            currency_suffix = f" to {budget.actual_currency}" if budget.actual_currency else ""
            ws.cell(row=plan["footnote_row"], column=1, value="*").font = _FOOTNOTE_STAR_FONT
            if estimated:
                text = (
                    "Amount and exchange rate estimated using the budget's estimated "
                    f"exchange rate — {estimated} of {len(rows)} expense row(s) not yet "
                    f"linked to a currency conversion{currency_suffix}."
                )
            else:
                omitted = (
                    "Exchange rate"
                    if no_rate < len(rows)
                    else ("Converted amount and exchange rate")
                )
                text = (
                    f"{omitted} omitted — {no_rate} of {len(rows)} expense row(s) "
                    f"not yet converted{currency_suffix}."
                )
            text_cell = ws.cell(row=plan["footnote_row"], column=2, value=text)
            text_cell.font = _AUDIT_FONT

    def _write_footer(self, plan: dict) -> None:
        self._write_audit_footer(plan["audit_row"])
