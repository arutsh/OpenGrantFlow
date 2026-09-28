import io
import struct
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

import app.services.template_detection.workbook_limits as workbook_limits
from app.services.template_detection.workbook_limits import (
    MAX_ZIP_MEMBERS,
    InvalidWorkbookError,
    WorkbookTooLargeError,
    check_workbook_limits,
)

_MINIMAL_EXTRAS = {
    "[Content_Types].xml": "<Types/>",
    "_rels/.rels": "<Relationships/>",
    "xl/workbook.xml": "<workbook/>",
}


def _sheet_xml(cell_refs: list[str]) -> str:
    cells = "".join(f'<c r="{ref}"/>' for ref in cell_refs)
    return (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"<sheetData><row>{cells}</row></sheetData></worksheet>"
    )


def _build_zip(sheet_cell_refs: list[str] | None = None, extra_files: dict | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in (extra_files if extra_files is not None else _MINIMAL_EXTRAS).items():
            zf.writestr(name, content)
        if sheet_cell_refs is not None:
            zf.writestr("xl/worksheets/sheet1.xml", _sheet_xml(sheet_cell_refs))
    return buf.getvalue()


def _forge_declared_uncompressed_size(data: bytes, filename: bytes, forged_size: int) -> bytes:
    """Patches the local + central directory `file_size` fields for `filename`,
    leaving the real compressed bytes untouched — simulates a hostile zip bomb."""
    data = bytearray(data)
    forged = struct.pack("<I", forged_size)

    idx = data.find(b"PK\x03\x04")
    while idx != -1:
        name_len = struct.unpack_from("<H", data, idx + 26)[0]
        name_start = idx + 30
        if bytes(data[name_start:name_start + name_len]) == filename:
            data[idx + 22:idx + 26] = forged
        idx = data.find(b"PK\x03\x04", idx + 4)

    idx = data.find(b"PK\x01\x02")
    while idx != -1:
        name_len = struct.unpack_from("<H", data, idx + 28)[0]
        name_start = idx + 46
        if bytes(data[name_start:name_start + name_len]) == filename:
            data[idx + 24:idx + 28] = forged
        idx = data.find(b"PK\x01\x02", idx + 4)

    return bytes(data)


def _real_workbook_bytes() -> bytes:
    wb = Workbook()
    wb.active.append(["Category", "Description", 100])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _append_member(data: bytes, name: str, content: bytes) -> bytes:
    buf = io.BytesIO(data)
    with zipfile.ZipFile(buf, "a") as zf:
        zf.writestr(name, content)
    return buf.getvalue()


class TestCheckWorkbookLimits:
    @pytest.mark.parametrize("target", [
        "/xl/worksheets/custom.xml",
        "worksheets/custom.xml",
        "custom.xml",
        "../xl/custom.xml",
    ])
    @pytest.mark.parametrize("row", [1, 10001])
    def test_relationship_targets_are_scanned(self, target, row):
        wb = Workbook()
        wb.active.cell(row=row, column=1, value="Budget")
        original = io.BytesIO()
        wb.save(original)
        path = "xl/worksheets/custom.xml" if "worksheets/" in target else "xl/custom.xml"
        rewritten = io.BytesIO()
        with (
            zipfile.ZipFile(original) as source,
            zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as dest,
        ):
            for name in source.namelist():
                content = source.read(name)
                if name == "xl/_rels/workbook.xml.rels":
                    content = content.replace(b"/xl/worksheets/sheet1.xml", target.encode())
                elif name == "[Content_Types].xml":
                    content = content.replace(b"/xl/worksheets/sheet1.xml", f"/{path}".encode())
                elif name == "xl/worksheets/sheet1.xml":
                    name = path
                dest.writestr(name, content)
        data = rewritten.getvalue()
        assert load_workbook(io.BytesIO(data)).active.max_row == row
        if row > workbook_limits.MAX_ROWS:
            with pytest.raises(WorkbookTooLargeError) as exc:
                check_workbook_limits(data)
            assert exc.value.reason == "max_rows"
        else:
            check_workbook_limits(data)

    @pytest.mark.parametrize("external", [False, True])
    def test_invalid_worksheet_relationship_is_rejected(self, external):
        extras = dict(_MINIMAL_EXTRAS)
        mode = ' TargetMode="External"' if external else ""
        extras["xl/_rels/workbook.xml.rels"] = (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            f'Target="missing.xml"{mode}/></Relationships>'
        )
        with pytest.raises(InvalidWorkbookError) as exc:
            check_workbook_limits(_build_zip(extra_files=extras))
        assert exc.value.reason == "invalid_workbook"

    def test_normal_template_passes(self):
        data = _build_zip(["A1", "B1", "C2"])
        check_workbook_limits(data)

    def test_sparse_huge_used_range_sheet_is_rejected(self):
        """A1 + XFD1048576 is a ~17 billion cell used range if openpyxl ever
        materializes it, despite the file itself being a few hundred bytes."""
        data = _build_zip(["A1", "XFD1048576"])
        with pytest.raises(WorkbookTooLargeError) as exc:
            check_workbook_limits(data)
        assert exc.value.reason == "max_rows"

    def test_too_many_zip_members_is_rejected(self):
        extras = {f"file{i}.xml": "x" for i in range(MAX_ZIP_MEMBERS + 1)}
        data = _build_zip(sheet_cell_refs=None, extra_files=extras)
        with pytest.raises(WorkbookTooLargeError) as exc:
            check_workbook_limits(data)
        assert exc.value.reason == "max_zip_members"

    def test_high_compression_ratio_member_is_rejected(self, monkeypatch):
        monkeypatch.setattr(workbook_limits, "RATIO_CHECK_THRESHOLD", 100)
        monkeypatch.setattr(workbook_limits, "MAX_COMPRESSION_RATIO", 10)
        payload = "A" * 10_000  # trivially compressible, ratio far over 10:1
        data = _build_zip(sheet_cell_refs=None, extra_files={"xl/bomb.xml": payload})
        with pytest.raises(WorkbookTooLargeError) as exc:
            check_workbook_limits(data)
        assert exc.value.reason == "max_compression_ratio"

    def test_forged_declared_size_does_not_bypass_the_cap(self, monkeypatch):
        """The zip's own `file_size` header can be forged smaller than the real
        content — the guard must decompress and count for itself, not trust it."""
        monkeypatch.setattr(workbook_limits, "MAX_MEMBER_SIZE", 1000)
        real_payload = "A" * 5000
        data = _build_zip(sheet_cell_refs=None, extra_files={"xl/bomb.xml": real_payload})
        forged = _forge_declared_uncompressed_size(data, b"xl/bomb.xml", forged_size=1)

        with zipfile.ZipFile(io.BytesIO(forged)) as zf:
            assert zf.infolist()[0].file_size == 1  # confirms the forgery took effect

        with pytest.raises(WorkbookTooLargeError) as exc:
            check_workbook_limits(forged)
        assert exc.value.reason == "max_member_size"

    def test_real_workbook_with_no_extra_members_passes(self):
        check_workbook_limits(_real_workbook_bytes())

    def test_polyglot_with_extra_member_is_rejected(self):
        """A genuinely openpyxl-loadable xlsx that also smuggles an unrelated
        file in the same zip container (see design.md Decision 5)."""
        data = _append_member(_real_workbook_bytes(), "payload.exe", b"MZ fake-binary-content")

        with pytest.raises(InvalidWorkbookError) as exc:
            check_workbook_limits(data)
        assert exc.value.reason == "unexpected_member"

    @pytest.mark.parametrize(
        "member",
        ["customXml/item1.xml", "customXml/_rels/item1.xml.rels", "docMetadata/LabelInfo.xml"],
    )
    def test_office_metadata_members_pass(self, member):
        """SharePoint/OneDrive customXml and Purview sensitivity-label parts."""
        data = _append_member(_real_workbook_bytes(), member, b"<?xml version='1.0'?><x/>")

        check_workbook_limits(data)


def _workbook_with_sheet_content(rows: str, merges: str = "") -> bytes:
    sheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{rows}</sheetData><mergeCells>{merges}</mergeCells></worksheet>'
    )
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(_real_workbook_bytes())) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            dest.writestr(name, sheet if name == "xl/worksheets/sheet1.xml" else source.read(name))
    return output.getvalue()


@pytest.mark.parametrize("rows, limit", [
    ('<row r="10001"><c><v>1</v></c></row>', "max_rows"),
    ('<row r="10000"><c/></row><row><c/></row>', "max_rows"),
    ('<row><c r="SF1"/><c/></row>', "max_columns"),
    ('<row>' + '<c/>' * 501 + '</row>', "max_columns"),
])
def test_implicit_coordinates_exceeding_limits_are_rejected(rows, limit):
    data = _workbook_with_sheet_content(rows)
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.max_row > workbook_limits.MAX_ROWS or ws.max_column > workbook_limits.MAX_COLUMNS
    with pytest.raises(WorkbookTooLargeError) as exc:
        check_workbook_limits(data)
    assert exc.value.reason == limit


@pytest.mark.parametrize("merges, limit", [
    ('<mergeCell ref="A1:A10001"/>', "max_rows"),
    ('<mergeCell ref="A1:SG1"/>', "max_columns"),
    ('<mergeCell ref="A1:SF2001"/>', "max_cells"),
    ('<mergeCell ref="A1:SF1001"/><mergeCell ref="A1002:SF2002"/>', "max_cells"),
])
def test_merged_ranges_exceeding_limits_are_rejected(merges, limit):
    data = _workbook_with_sheet_content('<row><c/></row>', merges)
    with pytest.raises(WorkbookTooLargeError) as exc:
        check_workbook_limits(data)
    assert exc.value.reason == limit


def test_small_implicit_cells_and_merges_pass():
    data = _workbook_with_sheet_content(
        '<row r="2"><c r="B2"><v>1</v></c><c><v>2</v></c></row>'
        '<row><c><v>3</v></c></row>',
        '<mergeCell ref="D1:E2"/>',
    )
    check_workbook_limits(data)
    ws = load_workbook(io.BytesIO(data)).active
    assert ws['B2'].value == 1
    assert ws['C2'].value == 2
    assert ws['A3'].value == 3
    assert str(next(iter(ws.merged_cells))) == 'D1:E2'


@pytest.mark.parametrize("rows, merges", [
    ('<row><c r="bogus"/></row>', ''),
    ('<row><c r="A0"/></row>', ''),
    ('<row r="1.5"><c/></row>', ''),
    ('', '<mergeCell ref="B2:A1"/>'),
    ('', '<mergeCell ref="A:A"/>'),
    ('', '<mergeCell/>'),
])
def test_malformed_coordinates_are_rejected(rows, merges):
    with pytest.raises(InvalidWorkbookError) as exc:
        check_workbook_limits(_workbook_with_sheet_content(rows, merges))
    assert exc.value.reason == "invalid_workbook"


def test_trailing_empty_rows_beyond_row_cap_are_not_rejected():
    """LibreOffice sometimes emits trailing empty `<row>` stubs (height/format
    metadata only) up to the sheet's max row number — harmless, no cells."""
    data = _workbook_with_sheet_content(
        '<row r="1"><c r="A1"><v>1</v></c></row>'
        '<row r="1048571"></row><row r="1048576"></row>'
    )
    check_workbook_limits(data)


def test_implicit_cells_count_toward_cell_limit(monkeypatch):
    monkeypatch.setattr(workbook_limits, "MAX_CELLS", 2)
    data = _workbook_with_sheet_content('<row><c/><c/><c/></row>')
    with pytest.raises(WorkbookTooLargeError) as exc:
        check_workbook_limits(data)
    assert exc.value.reason == "max_cells"
